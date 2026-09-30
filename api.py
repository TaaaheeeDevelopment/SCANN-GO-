import io
import os
import hmac
import hashlib
from datetime import datetime, timedelta, timezone
from functools import wraps

import pymysql
from flask import Flask, request, jsonify, session, send_file
from openpyxl import Workbook

SECRET_KEY = os.environ["SECRET_KEY"]
ITERATIONS = 600_000
BAGHDAD = timezone(timedelta(hours=3))
DAYS_AR = ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]

app = Flask(__name__)
app.config.update(
    SECRET_KEY=SECRET_KEY,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Strict",
    SESSION_COOKIE_SECURE=os.environ.get("COOKIE_SECURE", "1") == "1",
)
ADMINS = []  # [{"username": ..., "password_hash": ...}, ...]


# ---------- الهاش ----------
def hash_password(password):
    salt = os.urandom(16)
    dk = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, ITERATIONS)
    return f"{ITERATIONS}${salt.hex()}${dk.hex()}"


def check_password(password, stored):
    iterations, salt_hex, hash_hex = stored.split("$")
    dk = hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(salt_hex), int(iterations))
    return hmac.compare_digest(dk.hex(), hash_hex)


DUMMY_HASH = hash_password("dummy")


# ---------- الـ DB ----------
def get_connection():
    return pymysql.connect(
        host=os.environ.get("DB_HOST", "localhost"),
        user=os.environ["DB_USER"],
        password=os.environ["DB_PASSWORD"],
        database=os.environ.get("DB_NAME", "attendance_db"),
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
        autocommit=True,
    )


def query(sql, params=(), one=False):
    """SELECT: ترجع صف واحد أو كل الصفوف."""
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cur.execute(sql, params)
            return cur.fetchone() if one else cur.fetchall()
    finally:
        conn.close()


def execute(sql, params=()):
    """INSERT/UPDATE: ترجع (عدد الصفوف, آخر id)."""
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cur.execute(sql, params)
            return cur.rowcount, cur.lastrowid
    finally:
        conn.close()


# ---------- الوقت ----------
def now_baghdad():
    return datetime.now(BAGHDAD).replace(tzinfo=None)


def day_name(dt):
    return DAYS_AR[dt.weekday()]


def day_range(d):
    start = datetime(d.year, d.month, d.day)
    return start, start + timedelta(days=1)


def parse_date(value, default):
    if not value:
        return default
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        return None


# ---------- الأدمنات ----------
def load_admins():
    """يسحب اسم الطالب والهاش بعمل ربط JOIN بين جدول الأدمن وجدول الطلاب"""
    global ADMINS
    sql = """
        SELECT 
            ss.student_name AS username, 
            sa.password_hash,
            sa.student_id
        FROM scanner_admins sa
        JOIN section_student ss ON sa.student_id = ss.student_id
    """
    ADMINS = [dict(r) for r in query(sql)]
    return len(ADMINS)


def add_admin(username, password):
    execute("INSERT INTO admins (username, password_hash) VALUES (%s, %s)",
            (username, hash_password(password)))
    load_admins()


def find_admin(username):
    return next((a for a in ADMINS if a["username"] == username), None)


def verify_login(username, password):
    admin = find_admin(username)
    ok = check_password(password, admin["password_hash"] if admin else DUMMY_HASH)
    return bool(admin) and ok


def login_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if "admin" not in session:
            return jsonify({"ok": False, "error": "غير مصرح"}), 401
        return fn(*args, **kwargs)
    return wrapper


# ---------- الطلاب و QR ----------
def add_student(name, section):
    return execute("INSERT INTO section_student (student_name, section_name) VALUES (%s, %s)",
                   (name, section))[1]


def get_students_by_section(section):
    """تجيب كل طلاب الشعبة المحددة."""
    return query("SELECT student_id, student_name, section_name FROM section_student "
                 "WHERE section_name = %s ORDER BY student_name", (section,))


def find_student(name, section):
    return query("SELECT student_id, student_name, section_name FROM section_student "
                 "WHERE student_name = %s AND section_name = %s LIMIT 2", (name, section))


def make_qr_token(name, section):
    """الكلمة الي تنحط بالـ QR: الاسم|الشعبة"""
    return f"{name}|{section}"


def parse_qr_token(token):
    """ترجع (الاسم, الشعبة) أو None."""
    if "|" not in token or len(token) > 250:
        return None
    name, section = token.rsplit("|", 1)
    name, section = name.strip(), section.strip()
    return (name, section) if name and section else None


# ---------- الدكاترة والاستثناءات ----------
def find_doctor_id(name):
    row = query("SELECT dactor_id FROM dactor WHERE dactor_name = %s", (name,), one=True)
    return row["dactor_id"] if row else None


def get_or_create_doctor(name):
    return find_doctor_id(name) or execute(
        "INSERT INTO dactor (dactor_name) VALUES (%s)", (name,))[1]


def has_exception(student_id, lecture_section):
    """هل الطالب مسموح له يحضر بشعبة ثانية (جدول student_exceptions)؟"""
    row = query("SELECT 1 AS ok FROM student_exceptions "
                "WHERE student_id = %s AND allowed_section = %s LIMIT 1",
                (student_id, lecture_section), one=True)
    return row is not None


# ---------- الحضور ----------
def mark_present(student_id, dactor_id, when):
    """تسجل الحضور بوقت السيرفر. ترجع False إذا مسجل اليوم ويا نفس الدكتور."""
    start, end = day_range(when.date())
    rowcount, _ = execute("""
        INSERT INTO attendance_section (student_id, attendance_time, dactor_id)
        SELECT %s, %s, %s FROM DUAL
        WHERE NOT EXISTS (
            SELECT 1 FROM attendance_section
            WHERE student_id = %s AND dactor_id = %s
              AND attendance_time >= %s AND attendance_time < %s
        )
    """, (student_id, when, dactor_id, student_id, dactor_id, start, end))
    return rowcount == 1


def record_attendance(token, lecture_section, doctor_name):
    """دالة الحضور: تقرأ محتوى QR (اسم|شعبة)، تتحقق من الـ DB، تسجل، وترجع اليوم والتاريخ."""
    parsed = parse_qr_token(token)
    if not parsed:
        return {"ok": False, "error": "QR غير صالح"}, 400
    name, section = parsed

    if not lecture_section or len(lecture_section) > 100:
        return {"ok": False, "error": "الشعبة مطلوبة"}, 400

    matches = find_student(name, section)
    if not matches:
        return {"ok": False, "error": "الطالب غير موجود أو الشعبة غير مطابقة"}, 404
    if len(matches) > 1:
        return {"ok": False, "error": "اسم مكرر بنفس الشعبة"}, 409
    student = matches[0]

    is_guest = student["section_name"] != lecture_section
    if is_guest and not has_exception(student["student_id"], lecture_section):
        return {"ok": False, "error": "الطالب غير مسموح له بالحضور بهذي الشعبة"}, 403

    when = now_baghdad()
    doctor_id = get_or_create_doctor(doctor_name)
    if not mark_present(student["student_id"], doctor_id, when):
        return {"ok": False, "error": f"{name} مسجل مسبقاً اليوم"}, 409

    return {
        "ok": True,
        "name": student["student_name"],
        "section": student["section_name"],
        "is_guest": is_guest,
        "day": day_name(when),
        "date": when.strftime("%Y-%m-%d"),
        "time": when.strftime("%H:%M"),
    }, 200


def get_absent_students(section, dactor_id, date):
    start, end = day_range(date)
    return query("""
        SELECT student_id, student_name FROM section_student
        WHERE section_name = %s AND student_id NOT IN (
            SELECT student_id FROM attendance_section
            WHERE dactor_id = %s AND attendance_time >= %s AND attendance_time < %s)
        ORDER BY student_name
    """, (section, dactor_id, start, end))


# ---------- تصدير Excel ----------
def get_attendance_rows(date_from, date_to, section=None):
    """تجيب جدول الحضور من الـ DB ضمن فترة."""
    start, end = day_range(date_from)[0], day_range(date_to)[1]
    sql = """
        SELECT s.student_name, s.section_name, d.dactor_name, a.attendance_time
        FROM attendance_section a
        JOIN section_student s ON s.student_id = a.student_id
        JOIN dactor d ON d.dactor_id = a.dactor_id
        WHERE a.attendance_time >= %s AND a.attendance_time < %s
    """
    params = [start, end]
    if section:
        sql += " AND s.section_name = %s"
        params.append(section)
    return query(sql + " ORDER BY a.attendance_time", params)


def write_row(ws, values):
    ws.append(values)
    for cell in ws[ws.max_row]:
        if cell.data_type == "f":      # يمنع تنفيذ أي نص يبدأ بـ = كمعادلة
            cell.data_type = "s"


def build_excel(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "الحضور"
    ws.sheet_view.rightToLeft = True
    write_row(ws, ["#", "اسم الطالب", "الشعبة", "الدكتور", "اليوم", "التاريخ", "الوقت"])
    for i, r in enumerate(rows, 1):
        t = r["attendance_time"]
        write_row(ws, [i, r["student_name"], r["section_name"], r["dactor_name"],
                       day_name(t), t.strftime("%Y-%m-%d"), t.strftime("%H:%M")])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------- الـ Routes ----------
@app.post("/api/auth/login")
def login():
    data = request.get_json(silent=True) or {}
    username = str(data.get("username", ""))
    if verify_login(username, str(data.get("password", ""))):
        session.clear()
        session["admin"] = username
        return jsonify({"ok": True})
    return jsonify({"ok": False, "error": "بيانات غير صحيحة"}), 401


@app.get("/api/students")
@login_required
def students_route():
    section = request.args.get("section", "").strip()
    if not section:
        return jsonify({"ok": False, "error": "الشعبة مطلوبة"}), 400
    return jsonify(get_students_by_section(section))


@app.post("/api/attendance/scan")
@login_required
def scan_route():
    data = request.get_json(silent=True) or {}
    payload, status = record_attendance(
        str(data.get("token", "")),
        str(data.get("lecture_section", "")).strip(),
        str(data.get("doctor", "")).strip()[:100] or "د. المحاضر",
    )
    return jsonify(payload), status


@app.get("/api/attendance/absent")
@login_required
def absent_route():
    date = parse_date(request.args.get("date"), now_baghdad().date())
    doctor_id = find_doctor_id(request.args.get("doctor", "").strip())
    section = request.args.get("section", "").strip()
    if not date or not doctor_id or not section:
        return jsonify({"ok": False, "error": "معطيات ناقصة"}), 400
    return jsonify(get_absent_students(section, doctor_id, date))


@app.get("/api/attendance/export")
@login_required
def export_route():
    today = now_baghdad().date()
    date_from = parse_date(request.args.get("from"), today)
    date_to = parse_date(request.args.get("to"), today)
    if not date_from or not date_to or date_from > date_to:
        return jsonify({"ok": False, "error": "تاريخ غير صالح"}), 400
    section = request.args.get("section", "").strip() or None
    buf = build_excel(get_attendance_rows(date_from, date_to, section))
    return send_file(
        buf, as_attachment=True,
        download_name=f"attendance_{date_from}_{date_to}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


load_admins()   # أول شي يشتغل (برا __main__ لأن gunicorn ما يشغّله)

if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000)